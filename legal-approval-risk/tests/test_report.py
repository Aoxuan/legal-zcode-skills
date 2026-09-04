# -*- coding: utf-8 -*-
"""风险报告生成:命中风险项整行黄色高亮(2026-09-04 口径)。不发起网络请求。"""
import os
import sys
import tempfile
import unittest

_SCRIPTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts")
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)

from openpyxl import load_workbook

from qcc.report import build_report


def _payloads(status="存续", verify=1, shixin=None):
    return {
        2006: {"Result": {"VerifyResult": 1, "Data": {
            "Name": "甲公司", "CreditCode": "91ABC", "OperName": "张三",
            "Status": status, "ShiXin": shixin, "Exception": None,
        }}},
        856: {"Status": "200", "Result": {"VerifyResult": verify}},
    }


def _row_by_name(ws, name):
    for row in ws.iter_rows(min_row=1, max_col=2):
        if row[0].value == name:
            return row
    return None


def _is_yellow(cell):
    fill = cell.fill
    return (fill is not None and fill.patternType == "solid"
            and str(getattr(fill.start_color, "rgb", "") or "").endswith("FFFF00"))


class ReportHighlightTests(unittest.TestCase):
    def _build(self, payloads):
        with tempfile.TemporaryDirectory() as tmp:
            path = build_report("甲公司", payloads, tmp,
                                profile={"query_date": "2026-09-04 00:00:00",
                                         "risk_level": "high", "hits": [1],
                                         "called_apis": [2006, 856], "cost": 0})
            return load_workbook(path).active

    def test_hit_row_highlighted(self):
        ws = self._build(_payloads(shixin=[{"case": 1}]))
        row = _row_by_name(ws, "失信被执行人")
        self.assertIsNotNone(row)
        self.assertEqual(row[1].value, "1条")
        self.assertTrue(_is_yellow(row[0]))
        self.assertTrue(_is_yellow(row[1]))

    def test_no_record_row_not_highlighted(self):
        ws = self._build(_payloads())
        row = _row_by_name(ws, "经营异常名录")
        self.assertEqual(row[1].value, "无记录")
        self.assertFalse(_is_yellow(row[0]))

    def test_subject_rows_never_highlighted(self):
        ws = self._build(_payloads())
        row = _row_by_name(ws, "企业全称")
        self.assertEqual(row[1].value, "甲公司")
        self.assertFalse(_is_yellow(row[0]))

    def test_normal_status_not_highlighted(self):
        ws = self._build(_payloads(status="存续"))
        row = _row_by_name(ws, "存续/注销/吊销/迁出")
        self.assertEqual(row[1].value, "存续")
        self.assertFalse(_is_yellow(row[0]))

    def test_revoked_status_highlighted(self):
        ws = self._build(_payloads(status="吊销"))
        row = _row_by_name(ws, "存续/注销/吊销/迁出")
        self.assertEqual(row[1].value, "吊销")
        self.assertTrue(_is_yellow(row[0]))

    def test_three_element_mismatch_highlighted(self):
        ws = self._build(_payloads(verify=2))
        row = _row_by_name(ws, "三要素(代码+名称+法人)一致性")
        self.assertEqual(row[1].value, "企业名称不一致(2)")
        self.assertTrue(_is_yellow(row[0]))

    def test_three_element_match_not_highlighted(self):
        ws = self._build(_payloads(verify=1))
        row = _row_by_name(ws, "三要素(代码+名称+法人)一致性")
        self.assertEqual(row[1].value, "一致(VerifyResult=1)")
        self.assertFalse(_is_yellow(row[0]))

    def test_legend_row_present(self):
        ws = self._build(_payloads())
        row = _row_by_name(ws, "高亮说明")
        self.assertIsNotNone(row)
        self.assertIn("黄色行=命中风险项", row[1].value)


if __name__ == "__main__":
    unittest.main()
