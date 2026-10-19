# -*- coding: utf-8 -*-
"""合同主体一致性校验(2026-10-19 口径):docx 抬头解析、方向规则、
mismatch→reject / skip→放行提醒 / match→原风险流程。不发起网络请求。"""
import io
import os
import sys
import tempfile
import unittest
from unittest import mock

_SCRIPTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts")
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)

from common.company_store import MemoryCompanyStore
from common.contract_subject import (
    ContractParseError, clean_subject_value, looks_like_company, names_match,
    normalize_name, parse_contract_subjects, resolve_counterparty,
)
from common.state import StateStore
from poll import _process_one

_FIELDS = ["客户名称", "供应商名称", "达人企业名称"]
_TEST_ID = "11010519491231002X"
_OUR = "成都白兔有你文化传播有限公司"
_OUR2 = "成都抖咖文化传播有限公司"
_SUPPLIER = "杭州四叶草供应链管理有限公司"
_OTHER = "宁波樟树贸易有限公司"
_STAGE_NAME = "阿min学姐"  # 达人花名:无公司关键词,排除法应视为自然人排除


def _docx_blob(lines):
    from docx import Document
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _form(attach_value, field="供应商名称", company=_SUPPLIER):
    return [
        {"id": "cmp", "name": field, "type": "radioV2",
         "value": "%s｜91ABCD1234567890｜李四" % company},
        {"id": "our", "name": "签约主体", "type": "radioV2", "value": _OUR},
        {"id": "att", "name": "合同附件(请勿上传PDF、图片)", "type": "attachmentV2",
         "value": attach_value},
        {"id": "r1", "name": "企业风险评级", "type": "input", "value": ""},
        {"id": "a1", "name": "附件", "type": "attachmentV2", "value": ["https://old"]},
    ]


def _cfg(on=True):
    return {"field_companies": _FIELDS, "field_rating": "企业风险评级",
            "field_attachment": "附件", "approver_user_id": "ou_test",
            "approver_user_id_type": "open_id", "cache_days": 90,
            "field_contract_attach": "合同附件", "field_our_entity": "签约主体",
            "subject_check_enabled": on}


def _profile(level="low", source="cache"):
    return {"name": _SUPPLIER, "credit_code": "91ABCD1234567890", "legal_person": "李四",
            "status": "存续", "three_element_ok": True, "risk_level": level,
            "hits": [], "cost": 0, "source": source, "report_file": ""}


def _task(ic):
    return {"task_id": "t1", "approval_code": "AC1", "instance_code": ic}


class NormalizeTests(unittest.TestCase):
    def test_fullwidth_normalize(self):
        self.assertEqual(normalize_name("Ａ公司（上海）"), normalize_name("A公司(上海)"))

    def test_match_containment(self):
        self.assertTrue(names_match("杭州四叶草供应链管理有限公司",
                                    "杭州四叶草供应链管理有限公司分公司"))

    def test_short_names_not_matched_by_containment(self):
        # 过短名称禁止包含式误判
        self.assertFalse(names_match("甲公司", "甲公司群"))

    def test_mismatch(self):
        self.assertFalse(names_match(_SUPPLIER, _OTHER))


class ParseSubjectsTests(unittest.TestCase):
    def test_canonical_ab(self):
        blob = _docx_blob(["购销合同", "甲方：%s" % _OUR, "乙方：%s" % _SUPPLIER])
        self.assertEqual(parse_contract_subjects(blob),
                         {"party_a": _OUR, "party_b": _SUPPLIER})

    def test_alias_need_gong(self):
        blob = _docx_blob(["需方：%s" % _OUR, "供方：%s" % _SUPPLIER])
        self.assertEqual(parse_contract_subjects(blob),
                         {"party_a": _OUR, "party_b": _SUPPLIER})

    def test_qualifier_and_note(self):
        blob = _docx_blob(["甲方（买方）：%s（以下简称“甲方”）" % _OUR,
                           "乙方（卖方）：%s（以下简称“乙方”）" % _SUPPLIER])
        self.assertEqual(parse_contract_subjects(blob),
                         {"party_a": _OUR, "party_b": _SUPPLIER})

    def test_credit_code_suffix_cut(self):
        blob = _docx_blob(["甲方：%s" % _OUR,
                           "乙方：%s 统一社会信用代码：91330100MA2GL1234X" % _SUPPLIER])
        self.assertEqual(parse_contract_subjects(blob)["party_b"], _SUPPLIER)

    def test_placeholder_brackets_stripped(self):
        """模板占位括号【】:名称填在【】里,剥后取纯名称(2026-10-19 实测修订模式合同)。"""
        blob = _docx_blob(["甲方：【%s】" % _OTHER, "乙方：【%s】" % _OUR])
        subjects = parse_contract_subjects(blob)
        self.assertEqual(subjects, {"party_a": _OTHER, "party_b": _OUR})

    def test_empty_placeholder_side_missing(self):
        """未填的「【】」剥后为空 → 该侧缺失 → no_subject。"""
        blob = _docx_blob(["甲方：【】", "乙方：%s" % _OUR])
        with self.assertRaises(ContractParseError):
            parse_contract_subjects(blob)

    def test_one_line_dual_headings(self):
        """签署栏同行双抬头「甲方:【X】 乙方:Y」两侧都要解析到。"""
        blob = _docx_blob(["甲方：【%s】           乙方：%s" % (_OTHER, _OUR)])
        self.assertEqual(parse_contract_subjects(blob),
                         {"party_a": _OTHER, "party_b": _OUR})

    def test_three_parties_with_c(self):
        """达人合同三方:甲乙丙都解析,由排除法甄别。"""
        blob = _docx_blob(["甲方：%s" % _OUR2, "乙方：%s" % _SUPPLIER,
                           "丙方：%s" % _STAGE_NAME])
        self.assertEqual(parse_contract_subjects(blob),
                         {"party_a": _OUR2, "party_b": _SUPPLIER, "party_c": _STAGE_NAME})

    def test_paragraphs_inside_table(self):
        """签约方写在表格单元格里也要解析到(doc.paragraphs 不含表格,用 body 全量)。"""
        from docx import Document
        doc = Document()
        doc.add_paragraph("推广服务合同")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "甲方：%s" % _OUR
        table.cell(0, 1).text = "乙方：%s" % _SUPPLIER
        table.cell(1, 0).text = "账号：123"
        buf = io.BytesIO()
        doc.save(buf)
        self.assertEqual(parse_contract_subjects(buf.getvalue()),
                         {"party_a": _OUR, "party_b": _SUPPLIER})

    def test_not_docx_raises(self):
        with self.assertRaises(ContractParseError) as ctx:
            parse_contract_subjects(b"\x00\x01\x02\x03 unknown format")
        self.assertEqual(ctx.exception.code, "not_docx")

    def test_aligned_typeset_party_with_inner_spaces(self):
        """两端对齐排版把「丙方」拉成「丙    方」(2026-10-19 实测 202610190120
        漏丙方根因):角色词字符间容空白,三方都要解析到。"""
        blob = _docx_blob([
            "甲    方：【%s】" % _OUR2,
            "乙    方：【%s】" % _STAGE_NAME,
            "丙    方：【成都麦麦熊猫科技有限公司】",
        ])
        self.assertEqual(parse_contract_subjects(blob),
                         {"party_a": _OUR2, "party_b": _STAGE_NAME,
                          "party_c": "成都麦麦熊猫科技有限公司"})

    def test_legacy_doc_blob_utf16(self):
        """旧版 .doc(OLE2 魔数)走 UTF-16LE 扫描(202610190113 根因):
        「补充协议二甲方：X乙方：Y」前缀+连续双抬头都能解析。"""
        text = "补充协议二甲方：%s乙方：速奇办公设备（上海）有限公司    基于甲乙双方…" % _OUR
        blob = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + text.encode("utf-16-le"))
        subjects = parse_contract_subjects(blob)
        self.assertEqual(subjects.get("party_a"), _OUR)
        self.assertEqual(subjects.get("party_b"), "速奇办公设备（上海）有限公司")

    def test_pdf_branch_routes_and_wraps_errors(self):
        """%PDF 头进入 pdfplumber 分支;损坏 PDF 异常被包装成 parse_error。"""
        with self.assertRaises(ContractParseError) as ctx:
            parse_contract_subjects(b"%PDF-1.4 broken payload")
        self.assertEqual(ctx.exception.code, "parse_error")

    def test_pdf_no_text_layer_falls_to_no_subject(self):
        """扫描型 PDF(无文本层)提取为空 → no_subject → 调用方走人工核对兜底。"""
        from unittest import mock as _mock
        with _mock.patch("common.contract_subject._pdf_paragraphs", return_value=[]):
            with self.assertRaises(ContractParseError) as ctx:
                parse_contract_subjects(b"%PDF-1.4 fake")
        self.assertEqual(ctx.exception.code, "no_subject")

    def test_pdf_text_paragraphs_feed_same_pipeline(self):
        """文本型 PDF(Word 导出)提取行进入同一套抬头解析/排除法(口径:仅文本层 PDF)。"""
        from unittest import mock as _mock
        from common import contract_subject as cs
        lines = ["甲方：【%s】" % _OTHER, "乙方：%s" % _OUR]
        with _mock.patch.object(cs, "_pdf_paragraphs", return_value=lines):
            subjects = parse_contract_subjects(b"%PDF-1.4 fake")
        self.assertEqual(subjects, {"party_a": _OTHER, "party_b": _OUR})
        self.assertEqual(cs.resolve_counterparty(subjects, _OUR), ("company", _OTHER))

    def test_no_subject_raises(self):
        blob = _docx_blob(["这是一份没有主体抬头的说明文档"])
        with self.assertRaises(ContractParseError) as ctx:
            parse_contract_subjects(blob)
        self.assertEqual(ctx.exception.code, "no_subject")


class LooksLikeCompanyTests(unittest.TestCase):
    def test_chinese_company_keywords(self):
        for name in (_SUPPLIER, "某某设计工作室", "某某传媒中心", _OUR2):
            self.assertTrue(looks_like_company(name), name)

    def test_english_company_keywords(self):
        for name in ("Daketengwu Holding Limited", "Foo Pte. Ltd.", "Bar Inc."):
            self.assertTrue(looks_like_company(name), name)

    def test_natural_person_and_stage_name(self):
        """自然人姓名与达人花名(阿min学姐)不含公司关键词 → 判自然人排除。"""
        for name in ("王五", _STAGE_NAME, "林山清"):
            self.assertFalse(looks_like_company(name), name)


class ResolveCounterpartyTests(unittest.TestCase):
    """统一签约主体排除法(2026-10-19 二次纠偏):返回 (kind, name)。
    kind=company(唯一公司主体,比对用)/natural_person(合同对方是自然人而表单
    填了企业,形态不符)/unresolved(剩 0 或 ≥2 公司,人工核对兜底)。"""

    def test_exclusion_returns_other_side(self):
        subjects = {"party_a": _OUR, "party_b": _SUPPLIER}
        # 无论表单命中的是哪个字段,都按排除法取非我方一侧
        for field in ("供应商名称", "客户名称", "达人企业名称"):
            self.assertEqual(resolve_counterparty(subjects, _OUR, field),
                             ("company", _SUPPLIER))

    def test_counterparty_on_either_side(self):
        """合作方自行占甲方位:我方在乙方、对方在甲方也能取到。"""
        subjects = {"party_a": _OTHER, "party_b": _OUR}
        self.assertEqual(resolve_counterparty(subjects, _OUR, "供应商名称"),
                         ("company", _OTHER))

    def test_three_parties_exclude_our_and_person(self):
        """三方达人合同:排除我方+自然人(花名)后剩唯一公司主体。"""
        subjects = {"party_a": _OUR2, "party_b": _SUPPLIER, "party_c": _STAGE_NAME}
        self.assertEqual(resolve_counterparty(subjects, _OUR2, "达人企业名称"),
                         ("company", _SUPPLIER))

    def test_three_companies_unresolved(self):
        """排除我方后仍剩两家公司 → 无法唯一确定 → unresolved(走人工核对兜底)。"""
        subjects = {"party_a": _OUR, "party_b": _SUPPLIER, "party_c": _OTHER}
        self.assertEqual(resolve_counterparty(subjects, _OUR), ("unresolved", ""))

    def test_neither_matches_our_entity_unresolved(self):
        subjects = {"party_a": _OTHER, "party_b": "第三方法人主体有限公司"}
        self.assertEqual(resolve_counterparty(subjects, _OUR), ("unresolved", ""))

    def test_person_only_contract_is_natural_person(self):
        """只有我方+自然人(无对方公司):合同对方是个人而表单填了企业
        → 形态不符(2026-10-19 三次口径:拒绝,不再兜底放行)。"""
        subjects = {"party_a": _OUR2, "party_b": _STAGE_NAME}
        self.assertEqual(resolve_counterparty(subjects, _OUR2),
                         ("natural_person", _STAGE_NAME))

    def test_clean_subject_value(self):
        self.assertEqual(clean_subject_value("【%s】（以下简称“乙方”）。" % _SUPPLIER), _SUPPLIER)

    def test_clean_second_placeholder_segment(self):
        """并列第二占位段(实测「乙方:【赵馨】 昵称:【小福馨爱米线】」):
        剥整体包裹后截中间残余 】…【,取首段纯名称,防后段污染/误触公司关键词。"""
        self.assertEqual(clean_subject_value("【赵馨】 昵称：【小福馨爱米线】"), "赵馨")

    def test_clean_contact_info_suffix(self):
        """抬头行粘连联系人/联系方式段(实测「北京哆哆徕传媒有限公司联系人：吴家辉
        联系方式：18859638929」):按联系人类标记截断,取纯公司名。"""
        self.assertEqual(
            clean_subject_value("北京哆哆徕传媒有限公司联系人：吴家辉联系方式：18859638929"),
            "北京哆哆徕传媒有限公司")


class SubjectCheckProcessTests(unittest.TestCase):
    def _client(self, form, blob=None):
        client = mock.Mock()
        client.get_instance.return_value = {"status": "PENDING", "form_items": form}
        client.get_approval_form.return_value = []
        if blob is not None:
            client.download.return_value = blob
        return client

    def test_match_proceeds_risk_flow(self):
        blob = _docx_blob(["甲方：%s" % _OUR, "乙方：%s" % _SUPPLIER])
        client = self._client(_form(["https://contract.docx"]), blob)
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile", return_value=(_profile("low"), None)):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-MT"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["risk_level"], "low")
        self.assertEqual(rec["subject_check"]["status"], "match")
        self.assertEqual(rec["subject_check"]["contract_subject"], _SUPPLIER)
        client.approve_task.assert_called_once()
        client.reject_task.assert_not_called()

    def test_mismatch_rejects_without_qcc(self):
        blob = _docx_blob(["甲方：%s" % _OUR, "乙方：%s" % _OTHER])
        client = self._client(_form(["https://contract.docx"]), blob)
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile",
                            side_effect=AssertionError("主体不一致不应查风险")):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-MM"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertFalse(rec["approved"])
        self.assertTrue(rec["rejected"])
        self.assertEqual(rec["risk_level"], "subject_mismatch")
        self.assertEqual(rec["cost"], 0)
        client.reject_task.assert_called_once()
        comment = client.reject_task.call_args.kwargs["comment"]
        self.assertIn(_OTHER, comment)
        self.assertIn(_SUPPLIER, comment)
        self.assertIn("供应商名称", comment)
        self.assertIn("未查询企业风险", comment)
        client.approve_task.assert_not_called()
        client.upload_approval_file.assert_not_called()
        self.assertTrue(state.is_processed("INST-MM"))
        self.assertEqual(state.get("INST-MM")["reason"], "rejected_subject_mismatch")

    def test_mismatch_dry_run_neither_rejects_nor_marks(self):
        blob = _docx_blob(["甲方：%s" % _OUR, "乙方：%s" % _OTHER])
        client = self._client(_form(["https://contract.docx"]), blob)
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile",
                            side_effect=AssertionError("不应查风险")):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-MD"), cfg=_cfg(), dry_run=True)
        self.assertIsNone(err)
        self.assertFalse(rec["rejected"])
        client.reject_task.assert_not_called()
        self.assertFalse(state.is_processed("INST-MD"))

    def test_no_attachment_skips_to_manual_review(self):
        client = self._client(_form([]))
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile",
                            side_effect=AssertionError("未能校验不应查风险")):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-NA"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["rating_text"], "合同主体未能自动校验，请法务人工核对")
        self.assertEqual(rec["subject_check"]["status"], "skip")
        self.assertEqual(rec["cost"], 0)
        client.approve_task.assert_called_once()
        self.assertEqual(client.approve_task.call_args.kwargs["comment"],
                         "风险预审完成：合同主体未能自动校验，请法务人工核对，自动放行")
        client.reject_task.assert_not_called()
        client.upload_approval_file.assert_not_called()
        self.assertEqual(state.get("INST-NA")["reason"], "approved_subject_unverified")
        # 误传过风险报告附件(fixture 旧值)→清空
        form = client.approve_task.call_args.args[4]
        att = next(f for f in form if f["id"] == "a1")
        self.assertEqual(att["value"], [])

    def test_parse_fail_skips_to_manual_review(self):
        client = self._client(_form(["https://contract.docx"]), b"\xd0\xcf old doc")
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile",
                            side_effect=AssertionError("解析失败不应查风险")):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-PF"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["rating_text"], "合同主体未能自动校验，请法务人工核对")
        self.assertEqual(rec["subject_check"]["status"], "skip")

    def test_disabled_keeps_legacy_flow(self):
        blob = _docx_blob(["甲方：%s" % _OUR, "乙方：%s" % _OTHER])  # 即使不一致
        client = self._client(_form(["https://contract.docx"]), blob)
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile", return_value=(_profile("low"), None)):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-OFF"), cfg=_cfg(on=False), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        client.download.assert_not_called()
        client.reject_task.assert_not_called()
        self.assertEqual(state.get("INST-OFF")["reason"], "approved")

    def test_natural_person_skips_subject_check(self):
        store = MemoryCompanyStore()
        store.companies.append({"id": 1, "company_name": "王五", "credit_code": _TEST_ID,
                                "tax_no": _TEST_ID, "legal_person": ""})
        client = self._client(_form(["https://contract.docx"], field="供应商名称",
                                    company="王五"))
        profile = _profile()
        profile.update({"name": "王五", "risk_level": "natural_person", "source": "natural_person"})
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile", return_value=(profile, None)):
                rec, err = _process_one(client, store=store, state=state,
                                        task=_task("INST-NP"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["rating_text"], "自然人无须评估风险")
        client.download.assert_not_called()
        client.reject_task.assert_not_called()

    def test_counterparty_on_either_side_flow(self):
        """合作方占甲方位(我方在乙方):排除法照样取到甲方对方,一致→放行查风险。"""
        blob = _docx_blob(["甲方：%s" % _OTHER, "乙方：%s" % _OUR])
        client = self._client(_form(["https://contract.docx"], field="供应商名称",
                                    company=_OTHER), blob)
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile", return_value=(_profile("low"), None)):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-CS"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["subject_check"]["status"], "match")

    def test_real_contract_revision_placeholder_flow(self):
        """复刻真实修订模式合同(【合同】她叙&阿min学姐.docx 形态):
        甲方名称在【】里(修订插入)、抬头+签署栏同行双抬头;我方=成都抖咖。"""
        client_name = "杭州尚皙科技有限公司"
        blob = _docx_blob([
            "甲方：【%s】" % client_name, "联系地址：【】", "联系人：【】",
            "乙方：%s" % _OUR2, "联系人：【】",
            "……正文若干……",
            "甲方:【%s】           乙方:%s" % (client_name, _OUR2),
        ])
        form = _form(["https://contract.docx"], field="客户名称", company=client_name)
        # 我方签约主体=成都抖咖(这份真实合同的我方)
        for item in form:
            if item["name"] == "签约主体":
                item["value"] = _OUR2
        client = self._client(form, blob)
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile", return_value=(_profile("low"), None)):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-REV"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["subject_check"]["status"], "match")
        self.assertEqual(rec["subject_check"]["contract_subject"], client_name)

    def test_three_party_influencer_flow(self):
        """三方达人合同:甲=我方、乙=对方公司、丙=达人(花名无公司关键词)。
        排除我方+自然人后剩唯一公司=表单「达人企业名称」控件值 → match。"""
        blob = _docx_blob(["甲方：%s" % _OUR2, "乙方：%s" % _SUPPLIER,
                           "丙方：%s" % _STAGE_NAME])
        client = self._client(_form(["https://contract.docx"], field="达人企业名称",
                                    company=_SUPPLIER), blob)
        for item in client.get_instance.return_value["form_items"]:
            if item["name"] == "签约主体":
                item["value"] = _OUR2
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile", return_value=(_profile("low"), None)):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-3P"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["approved"])
        self.assertEqual(rec["subject_check"]["status"], "match")
        self.assertEqual(rec["subject_check"]["contract_subject"], _SUPPLIER)

    def test_three_party_aligned_typeset_mismatch_rejects(self):
        """202610190120 复刻:两端对齐排版三方(甲=我方抖咖/乙=自然人赵馨/
        丙=成都麦麦熊猫),表单「达人企业名称」却填成都讯格得 → 解析齐三方、
        排除法剩唯一公司丙方 → 比对不一致 → 拒绝。"""
        blob = _docx_blob([
            "甲    方：【%s】" % _OUR2,
            "乙    方：【赵馨】",
            "丙    方：【成都麦麦熊猫科技有限公司】",
        ])
        client = self._client(_form(["https://contract.docx"], field="达人企业名称",
                                    company="成都讯格得信息科技有限公司"), blob)
        for item in client.get_instance.return_value["form_items"]:
            if item["name"] == "签约主体":
                item["value"] = _OUR2
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile",
                            side_effect=AssertionError("mismatch 不应查风险")):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-3MM"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["rejected"])
        comment = client.reject_task.call_args.kwargs["comment"]
        self.assertIn("成都麦麦熊猫科技有限公司", comment)
        self.assertIn("成都讯格得信息科技有限公司", comment)
        self.assertEqual(state.get("INST-3MM")["reason"], "rejected_subject_mismatch")

    def test_person_counterparty_rejects(self):
        """合同对方仅剩自然人(表单填了企业):形态不符直接拒绝,评论带（自然人）标记
        (2026-10-19 三次口径,替代旧的兜底放行)。"""
        blob = _docx_blob(["甲方：%s" % _OUR2, "乙方：赵馨"])
        client = self._client(_form(["https://contract.docx"], field="达人企业名称",
                                    company="成都讯格得信息科技有限公司"), blob)
        for item in client.get_instance.return_value["form_items"]:
            if item["name"] == "签约主体":
                item["value"] = _OUR2
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            with mock.patch("poll.fetch_or_cache_profile",
                            side_effect=AssertionError("形态不符不应查风险")):
                rec, err = _process_one(client, store=MemoryCompanyStore(), state=state,
                                        task=_task("INST-NR"), cfg=_cfg(), dry_run=False)
        self.assertIsNone(err)
        self.assertTrue(rec["rejected"])
        comment = client.reject_task.call_args.kwargs["comment"]
        self.assertIn("「赵馨」（自然人）", comment)
        client.approve_task.assert_not_called()


class PollSkipProcessedTests(unittest.TestCase):
    """多节点流程幂等(2026-10-19):预审账号可能也在后续人工会签节点的待办里,
    预审已完成的实例(state 有记录)必须跳过,技能不代签人工节点。"""

    def test_poll_skips_processed_instances(self):
        import poll as poll_mod
        client = mock.Mock()
        client.query_user_tasks.return_value = [
            {"task_id": "t1", "approval_code": "AC1", "instance_code": "INST-DONE",
             "status": "PENDING", "support_api_operate": True},
            {"task_id": "t2", "approval_code": "AC1", "instance_code": "INST-NEW",
             "status": "PENDING", "support_api_operate": True},
        ]
        cfg = {"approval_codes": ["AC1"], "approver_user_id": "ou_test",
               "approver_user_id_type": "open_id", "comment_dry_run": False,
               "app_id": "ai", "app_secret": "as"}
        emitted = {}

        def _fake_process(client, store, state, task, cfg, dry_run, force=False):
            assert task["instance_code"] == "INST-NEW", "已处理实例不应再次处理"
            return {"instance_code": "INST-NEW", "approved": True}, None

        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            state.mark("INST-DONE", risk_level="low", reason="approved")
            with mock.patch("poll.check_all", return_value=(True, [])), \
                 mock.patch("poll._require_mysql_store", return_value=object()), \
                 mock.patch("poll.get_feishu", return_value=cfg), \
                 mock.patch("poll.FeishuClient", return_value=client), \
                 mock.patch("poll.StateStore", return_value=state), \
                 mock.patch("poll._process_one", side_effect=_fake_process), \
                 mock.patch("poll._emit", side_effect=lambda obj: emitted.update(obj)):
                rc = poll_mod.cmd_poll()
        self.assertEqual(rc, 0)
        self.assertEqual(emitted["skipped_processed"], ["INST-DONE"])
        self.assertEqual(emitted["new"], 1)
        self.assertEqual(emitted["ok"], True)

    def test_process_instance_already_processed_returns_note(self):
        import poll as poll_mod
        client = mock.Mock()
        client.query_user_tasks.return_value = [
            {"task_id": "t9", "approval_code": "AC1", "instance_code": "INST-DONE2",
             "status": "PENDING", "support_api_operate": True}]
        cfg = {"approver_user_id": "ou_test", "approver_user_id_type": "open_id",
               "app_id": "ai", "app_secret": "as"}
        emitted = {}
        with tempfile.TemporaryDirectory() as tmp:
            state = StateStore(os.path.join(tmp, "state.json"))
            state.mark("INST-DONE2", risk_level="low", reason="approved")
            with mock.patch("poll.check_all", return_value=(True, [])), \
                 mock.patch("poll._require_mysql_store", return_value=object()), \
                 mock.patch("poll.get_feishu", return_value=cfg), \
                 mock.patch("poll.FeishuClient", return_value=client), \
                 mock.patch("poll.StateStore", return_value=state), \
                 mock.patch("poll._process_one",
                            side_effect=AssertionError("不应再处理已预审实例")), \
                 mock.patch("poll._emit", side_effect=lambda obj: emitted.update(obj)):
                rc = poll_mod.cmd_process_instance("INST-DONE2")
        self.assertEqual(rc, 0)
        self.assertTrue(emitted["already_processed"])
        self.assertIn("人工审批节点", emitted["note"])


if __name__ == "__main__":
    unittest.main()
