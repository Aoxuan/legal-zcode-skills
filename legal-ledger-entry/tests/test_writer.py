# -*- coding: utf-8 -*-
"""schema 校验、主/子表写入、dry-run、幂等。离线 fake client。"""
import json
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from common.state import StateStore
from entry import run_entry
from ledger.errors import LedgerStop
from ledger.widgets import (
    APP_TOKEN,
    APPROVAL_CODE,
    CATEGORY_TABLE,
    MAIN_TABLE,
    SUB_TABLE,
    widgets_for,
)

# 表单 fixture 与 cfg 同源：都按 APPROVAL_CODE 解析控件 id，
# 常量指向正式/副本任一环境时测试都成立（supplier/risk_attach 随环境切换）。
W = widgets_for(APPROVAL_CODE)

CFG = {
    "app_id": "cli_test",
    "app_secret": "secret_test",
    "approval_code": APPROVAL_CODE,
    "app_token": APP_TOKEN,
    "main_table": MAIN_TABLE,
    "sub_table": SUB_TABLE,
    "category_table": CATEGORY_TABLE,
    "approver_user_id": "",
}


def _docx(title):
    doc = Document()
    run = doc.add_paragraph().add_run(title)
    run.font.size = Pt(22)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _form(our, **extra):
    items = [
        {"id": W["our_entity"], "value": our},
        {"id": W["contract_attach"], "value": ["https://file.example/contract.docx"]},
        {"id": W["risk_attach"], "value": ["https://risk.example/skip.docx"]},
    ]
    for key, value in extra.items():
        items.append({"id": W[key], "value": value})
    return items


class FakeClient:
    def __init__(self, form, docx, parents=None, main_options=None, category_ids=None):
        self.form = form
        self.docx = docx
        self.parents = list(parents or [])
        self.main_options = main_options or ["成都抖咖", "兔竞"]
        self.category_ids = list(category_ids or [
            "rec9hHL3YE", "rec5epcXPx", "recdDEDJNB", "recUejFCZo",
        ])
        self.created = []
        self.updated = []
        self.queries = []
        self.instance_reads = []
        self.serial = "202609090050"
        self.instance_code = "INST-1"

    def query_by_serial(self, approval_code, serial):
        self.queries.append((approval_code, serial))
        if serial == self.serial:
            return [{"instance_code": self.instance_code, "serial_number": serial}]
        return []

    def get_instance(self, instance_code):
        self.instance_reads.append(instance_code)
        return {"form": json.dumps(self.form, ensure_ascii=False), "instance_code": instance_code}

    def download(self, url):
        assert "risk" not in url
        return self.docx

    def bitable_list_fields(self, app_token, table_id):
        assert app_token == APP_TOKEN
        if table_id == MAIN_TABLE:
            options = self.main_options
            field_id = "fldXdDUQ0G"
        elif table_id == SUB_TABLE:
            options = ["成都兔竞", "成都抖咖"]
            field_id = "fldxX1CGKP"
        else:
            options = []
            field_id = ""
        return [{
            "field_id": field_id,
            "field_name": "我方主体",
            "property": {"options": [{"name": n} for n in options]},
        }]

    def bitable_list_records(self, app_token, table_id):
        assert table_id == CATEGORY_TABLE
        return [{"record_id": rid} for rid in self.category_ids]

    def bitable_search(self, app_token, table_id, field_name, value, operator="is"):
        assert table_id == MAIN_TABLE
        assert field_name == "合同名称唯一值"
        assert operator == "contains"
        return [p for p in self.parents if p.get("match", value) == value]

    def bitable_create_record(self, app_token, table_id, fields):
        self.created.append((table_id, fields))
        if table_id == SUB_TABLE:
            return "recNEW", {"子合同编号": "BTSP202307049-BC01-202610"}
        return "recNEW", {"合同编号（以此为准）": "BTSP202610099"}

    def bitable_get_record(self, app_token, table_id, record_id):
        raise AssertionError("create 响应已带回编号，不应轮询")

    def bitable_update_record(self, *args, **kwargs):
        self.updated.append((args, kwargs))
        raise AssertionError("不应更新记录")


def _run(fake, state, dry_run=False, serial="202609090050"):
    return run_entry(serial, dry_run=dry_run, client=fake, state=state, cfg=CFG)


def test_schema_entity_missing_stops(tmp_path):
    fake = FakeClient(
        _form("成都抖咖文化传播有限公司", customer="甲公司", amount="100", summary="概述"),
        _docx("直播带货合作协议"),
        main_options=["上海晏齐"],
    )
    state = StateStore(str(tmp_path / "state.json"))
    with pytest.raises(LedgerStop):
        _run(fake, state)
    assert fake.created == []
    assert fake.updated == []


def test_schema_category_missing_stops(tmp_path):
    fake = FakeClient(
        _form("成都兔竞贸易有限公司", customer="甲公司", amount="100", summary="概述"),
        _docx("广告补充协议BTSP202307049-BC01"),
        parents=[{"record_id": "recPARENT", "match": "BTSP202307049"}],
        category_ids=["recOTHER"],
    )
    state = StateStore(str(tmp_path / "state.json"))
    with pytest.raises(LedgerStop):
        _run(fake, state)
    assert fake.created == []


def test_write_main_and_read_back_number(tmp_path):
    fake = FakeClient(
        _form(
            "成都抖咖文化传播有限公司",
            customer="甲公司｜91ABC",
            amount="12000",
            summary="概述甲",
        ),
        _docx("直播带货合作协议"),
    )
    state = StateStore(str(tmp_path / "state.json"))
    result = _run(fake, state)
    assert result["ok"] is True
    assert result["contract_no"] == "BTSP202610099"
    assert result["record_id"] == "recNEW"
    assert fake.queries[0][0] == APPROVAL_CODE
    assert len(fake.created) == 1
    table_id, fields = fake.created[0]
    assert table_id == MAIN_TABLE
    assert fields["合同名称"] == "直播带货合作协议"
    assert fields["合同类型"] == "主合同"
    assert fields["我方主体"] == "成都抖咖"
    assert fields["相对方主体（全称）"] == "甲公司"
    assert fields["合同金额"] == "12000"
    assert fields["合同摘要（如：张喜喜&麦当劳23-6-6短视频）"] == "概述甲"
    for banned in (
        "合同编号（以此为准）", "合同编号BT/N", "合同名称唯一值",
        "子合同记录表", "合同属性", "签订日期",
    ):
        assert banned not in fields
    stored = state.get("INST-1")
    assert stored["record_id"] == "recNEW"
    assert stored["contract_no"] == "BTSP202610099"
    assert stored["table_id"] == MAIN_TABLE
    assert stored["kind"] == "主合同"


def test_write_sub_requires_unique_parent(tmp_path):
    fake = FakeClient(
        _form(
            "成都兔竞贸易有限公司",
            supplier="乙供应商｜CODE",
            amount="300",
            summary="补充概述",
        ),
        _docx("广告补充协议BTSP202307049-BC01"),
        parents=[{"record_id": "recPARENT", "match": "BTSP202307049"}],
    )
    state = StateStore(str(tmp_path / "state.json"))
    result = _run(fake, state)
    assert result["ok"] is True
    assert result["contract_no"] == "BTSP202307049-BC01-202610"
    table_id, fields = fake.created[0]
    assert table_id == SUB_TABLE
    assert "recPARENT" in fields["主合同"]
    assert "rec9hHL3YE" in fields["子合同类型"]
    assert fields["我方主体"] == "成都兔竞"
    assert fields["相对方主体（全称）"] == "乙供应商"
    assert fields["金额"] == "300"
    assert fields["摘要（概括子合同内容）"] == "补充概述"
    stored = state.get("INST-1")
    assert stored["table_id"] == SUB_TABLE
    assert stored["kind"] == "补充协议"
    assert stored["contract_no"] == "BTSP202307049-BC01-202610"


def test_sub_parent_not_unique_stops(tmp_path):
    fake = FakeClient(
        _form("成都抖咖文化传播有限公司", customer="甲公司", amount="1", summary="s"),
        _docx("广告补充协议BTSP202307049-BC01"),
        parents=[
            {"record_id": "recA", "match": "BTSP202307049"},
            {"record_id": "recB", "match": "BTSP202307049"},
        ],
    )
    state = StateStore(str(tmp_path / "state.json"))
    with pytest.raises(LedgerStop):
        _run(fake, state)
    assert fake.created == []


def test_dry_run_does_not_create_or_update(tmp_path):
    fake = FakeClient(
        _form(
            "成都抖咖文化传播有限公司",
            customer="甲公司",
            amount="9",
            summary="干跑摘要",
        ),
        _docx("直播带货合作协议"),
    )
    state = StateStore(str(tmp_path / "state.json"))
    result = _run(fake, state, dry_run=True)
    assert result["ok"] is True
    assert result["dry_run"] is True
    assert result["contract_no"] == "DRYRUN"
    assert result["release"] == "skipped"
    assert result["release_reason"] == "干跑不放行"
    assert result["table_id"] == MAIN_TABLE
    assert result["fields"]["合同名称"] == "直播带货合作协议"
    assert result["fields"]["合同类型"] == "主合同"
    assert fake.created == []
    assert fake.updated == []
    assert state.get("INST-1") is None
    assert fake.instance_reads == ["INST-1"]


def test_idempotent_when_state_has_record(tmp_path):
    fake = FakeClient(
        _form("成都抖咖文化传播有限公司", customer="甲公司", amount="1", summary="s"),
        _docx("直播带货合作协议"),
    )
    state = StateStore(str(tmp_path / "state.json"))
    state.mark(
        "INST-1",
        table_id=MAIN_TABLE,
        record_id="recOLD",
        contract_no="BTSP202600001",
        kind="主合同",
    )
    result = _run(fake, state)
    assert result["ok"] is True
    assert result["record_id"] == "recOLD"
    assert result["contract_no"] == "BTSP202600001"
    assert result.get("idempotent") is True
    assert fake.created == []
    assert fake.updated == []
    assert fake.instance_reads == []


def test_readback_api_error_still_marks_state(tmp_path):
    """建行后回读抛接口异常：state 仍记 record_id，重跑幂等返回已有行，不再建行。"""
    from feishu.client import FeishuApiError

    class ApiErrorReadClient(FakeClient):
        def bitable_create_record(self, app_token, table_id, fields):
            record_id, _fields = super().bitable_create_record(app_token, table_id, fields)
            return record_id, {}  # 创建响应不带编号，迫使走回读

        def bitable_get_record(self, app_token, table_id, record_id):
            raise FeishuApiError("回读台账记录失败")

    fake = ApiErrorReadClient(
        _form("成都抖咖文化传播有限公司", customer="甲公司", amount="1", summary="s"),
        _docx("直播带货合作协议"),
    )
    state = StateStore(str(tmp_path / "state.json"))
    with pytest.raises(FeishuApiError):
        _run(fake, state)
    stored = state.get("INST-1")
    assert stored is not None
    assert stored["record_id"] == "recNEW"
    assert "contract_no" not in stored

    class RecoveredClient(ApiErrorReadClient):
        def bitable_get_record(self, app_token, table_id, record_id):
            return {"合同编号（以此为准）": "BTSP202610099"}

    fake2 = RecoveredClient(
        _form("成都抖咖文化传播有限公司", customer="甲公司", amount="1", summary="s"),
        _docx("直播带货合作协议"),
    )
    result = _run(fake2, state)
    assert result["ok"] is True
    assert result.get("idempotent") is True
    assert result["record_id"] == "recNEW"
    assert result["contract_no"] == "BTSP202610099"
    assert fake2.created == []
    assert state.get("INST-1")["contract_no"] == "BTSP202610099"
