# -*- coding: utf-8 -*-
"""docx 标题、路由、父合同编号、旧 .doc 拒绝。离线，现造 docx。"""
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from ledger.docx_parse import assert_docx_bytes, extract_title, find_parent_no, route_kind
from ledger.errors import LedgerStop


def _docx(paragraphs):
    """paragraphs: [(text, pt|None), ...]"""
    doc = Document()
    for text, pt in paragraphs:
        run = doc.add_paragraph().add_run(text)
        if pt is not None:
            run.font.size = Pt(pt)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_title_is_largest_font_skipping_labels():
    data = _docx([
        ("甲方：", 28),
        ("直播带货合作协议", 22),
        ("正文小字说明", 12),
    ])
    assert extract_title(data) == "直播带货合作协议"


def test_title_skips_halfwidth_colon_label():
    data = _docx([
        ("Party:", 36),
        ("品牌服务合同书", 16),
    ])
    assert extract_title(data) == "品牌服务合同书"


def test_title_fallback_when_no_font_size():
    data = _docx([
        ("甲方：", None),
        ("合作协议正文标题", None),
        ("第二行说明", None),
    ])
    assert extract_title(data) == "合作协议正文标题"


@pytest.mark.parametrize("title,kind", [
    ("品牌补充协议", "补充协议"),
    ("年度补充合同", "补充协议"),
    ("终止协议书", "终止协议"),
    ("解除协议", "终止协议"),
    ("解约协议", "终止协议"),
    ("九月结算单", "结算单"),
    ("费用结算协议", "结算单"),
    ("结算对账单", "结算单"),
    ("投放执行单", "执行单"),
    ("执行确认单", "执行单"),
    ("直播带货合作协议", "主合同"),
])
def test_route_five_kinds(title, kind):
    assert route_kind(title) == kind


@pytest.mark.parametrize("title,body,expected", [
    ("补充协议BTSP202307049-BC01", "", "BTSP202307049"),
    ("补充协议BTSP202307049-ZZ02", "", "BTSP202307049"),
    ("补充协议BTSP202307049-JS3", "", "BTSP202307049"),
    ("执行单BTSP202307049-zx02", "", "BTSP202307049"),
    ("补充协议", "详见主合同BTSP2026094342条款", "BTSP2026094342"),
    ("BTSP202307049与BTSP202307050", "", None),
    ("补充协议", "没有父编号", None),
    ("主合同BTSP202307049", "", "BTSP202307049"),
    ("主合同BTSP2026094342", "", "BTSP2026094342"),
    ("BTSP202307049-BC01", "另见BTSP2026094342", "BTSP202307049"),
])
def test_parent_btsp(title, body, expected):
    assert find_parent_no(title, body) == expected


def test_old_doc_ole2_rejected():
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32
    with pytest.raises(LedgerStop):
        assert_docx_bytes(ole)
