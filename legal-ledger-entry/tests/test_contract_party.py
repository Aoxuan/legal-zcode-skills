# -*- coding: utf-8 -*-
"""合同签约内容兜底提取：抬头解析、清洗、排除法甄别。离线。"""
import io

from docx import Document

from ledger.contract_party import (
    clean_subject_value,
    looks_like_company,
    names_match,
    parse_contract_subjects,
    party_from_docx,
    resolve_counterparty,
)


def _docx(lines, table_rows=None):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell in enumerate(row):
                table.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_headings_in_body_and_table():
    blob = _docx(
        ["业务合作协议", "甲方：成都抖咖文化传播有限公司", "乙方：思搜网络（上海）有限公司"],
        table_rows=[["甲方：成都抖咖文化传播有限公司", "签署日期：2026年10月"]],
    )
    subjects = parse_contract_subjects(blob)
    assert subjects["party_a"] == "成都抖咖文化传播有限公司"
    assert subjects["party_b"] == "思搜网络（上海）有限公司"


def test_one_line_multi_heading_and_placeholder():
    blob = _docx(["甲方：【成都抖咖文化传播有限公司】 乙方：思搜网络（上海）有限公司（以下简称乙方）"])
    subjects = parse_contract_subjects(blob)
    assert subjects["party_a"] == "成都抖咖文化传播有限公司"
    assert subjects["party_b"] == "思搜网络（上海）有限公司"


def test_role_word_tolerates_inner_spaces():
    blob = _docx(["甲    方：成都抖咖文化传播有限公司", "乙　方：思搜网络（上海）有限公司"])
    subjects = parse_contract_subjects(blob)
    assert subjects["party_a"] == "成都抖咖文化传播有限公司"
    assert subjects["party_b"] == "思搜网络（上海）有限公司"


def test_clean_strips_credit_code_and_labels():
    assert clean_subject_value("思搜网络（上海）有限公司｜91310104MA1FR3NP9J｜刘川球") == \
        "思搜网络（上海）有限公司"
    assert clean_subject_value("甲公司 联系人：张三") == "甲公司"
    assert clean_subject_value("【】") == ""


def test_resolve_excludes_ours_and_picks_unique_company():
    blob = _docx(["甲方：成都抖咖文化传播有限公司", "乙方：思搜网络（上海）有限公司"])
    assert party_from_docx(blob, "成都抖咖文化传播有限公司") == \
        ("company", "思搜网络（上海）有限公司")


def test_ours_on_either_side_is_excluded():
    blob = _docx(["甲方：思搜网络（上海）有限公司", "乙方：成都抖咖文化传播有限公司"])
    assert party_from_docx(blob, "成都抖咖文化传播有限公司") == \
        ("company", "思搜网络（上海）有限公司")


def test_three_party_talent_contract_picks_company():
    blob = _docx(["甲方：某品牌有限公司", "乙方：成都抖咖文化传播有限公司", "丙方：张小美"])
    assert party_from_docx(blob, "成都抖咖文化传播有限公司") == ("company", "某品牌有限公司")


def test_natural_person_counterparty():
    blob = _docx(["甲方：成都抖咖文化传播有限公司", "乙方：李小美"])
    assert party_from_docx(blob, "成都抖咖文化传播有限公司") == ("natural_person", "李小美")


def test_unresolved_when_no_heading():
    blob = _docx(["双方经友好协商，达成如下条款"])
    assert party_from_docx(blob, "成都抖咖文化传播有限公司") == ("unresolved", "")


def test_unresolved_when_two_foreign_companies():
    blob = _docx(["甲方：甲公司", "乙方：乙公司", "丙方：成都抖咖文化传播有限公司"])
    assert party_from_docx(blob, "成都抖咖文化传播有限公司")[0] == "unresolved"


def test_names_match_contains_with_min_len():
    assert names_match("成都抖咖文化传播有限公司", "成都抖咖")
    assert not names_match("甲公司", "乙公司")


def test_looks_like_company():
    assert looks_like_company("思搜网络（上海）有限公司")
    assert looks_like_company("Daketengwu Holding Limited")
    assert not looks_like_company("阿min学姐")
